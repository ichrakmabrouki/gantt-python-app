"""
template_excel.py — génère le classeur modèle attendu par l'application.

Le parseur (`backend/solver/input_parser.py`) attend une structure très
précise : quatre feuilles aux noms exacts, émojis compris, et des en-têtes
placés sur des lignes déterminées. Personne ne peut deviner ça — d'où ce
générateur, branché sur un bouton de téléchargement dans l'écran d'import.

Le classeur produit contient un exemple complet et cohérent : 3 pièces,
6 opérations, 4 machines, 2 techniciens. L'utilisateur remplace les lignes
d'exemple par ses données.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Les noms de feuilles sont imposés par le parseur : ne pas les modifier sans
# modifier aussi input_parser.parse_excel_to_dict.
FEUILLE_PARAMETRES = "⚙ PARAMETRES"
FEUILLE_GAMMES = "📋 GAMMES"
FEUILLE_MODES = "🔧 MODES & PT"
FEUILLE_TECHNICIENS = "👷 TECHNICIENS"

_TITRE = Font(size=14, bold=True, color="1F3864")
_AIDE = Font(size=9, italic=True, color="7F7F7F")
_ENTETE = Font(bold=True, color="FFFFFF")
_FOND_ENTETE = PatternFill("solid", fgColor="1F3864")
_BORDURE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _ecrire_entete(ws, ligne: int, colonnes: list[str]) -> None:
    for i, nom in enumerate(colonnes, start=1):
        cellule = ws.cell(row=ligne, column=i, value=nom)
        cellule.font = _ENTETE
        cellule.fill = _FOND_ENTETE
        cellule.alignment = Alignment(horizontal="center", vertical="center")
        cellule.border = _BORDURE


def _ecrire_lignes(ws, depart: int, lignes: list[list]) -> None:
    for decalage, valeurs in enumerate(lignes):
        for i, valeur in enumerate(valeurs, start=1):
            cellule = ws.cell(row=depart + decalage, column=i, value=valeur)
            cellule.border = _BORDURE
            cellule.alignment = Alignment(horizontal="center")


def _largeurs(ws, largeurs: list[int]) -> None:
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur


def construire_template() -> bytes:
    """Retourne le classeur modèle sous forme d'octets, prêt à télécharger."""
    wb = Workbook()

    # ══ PARAMETRES ══ en-tête ligne 3, données à partir de la ligne 4 ═════════
    ws = wb.active
    ws.title = FEUILLE_PARAMETRES
    ws["A1"] = "Paramètres généraux de l'atelier"
    ws["A1"].font = _TITRE
    ws["A2"] = ("Renseignez la colonne Valeur. Ne déplacez pas la ligne "
                "d'en-tête : l'application la cherche en ligne 3.")
    ws["A2"].font = _AIDE
    _ecrire_entete(ws, 3, ["Paramètre", "Valeur", "Signification"])
    _ecrire_lignes(ws, 4, [
        ["nbJobs",  3,  "Nombre de pièces à ordonnancer"],
        ["nbMchs",  4,  "Nombre de machines de l'atelier"],
        ["nbOps",   6,  "Nombre total d'opérations (toutes pièces confondues)"],
        ["nbtechs", 2,  "Nombre de techniciens"],
        ["cte",     15, "Durée d'un changement de série, en minutes"],
    ])
    for ligne in range(4, 9):
        ws.cell(row=ligne, column=3).alignment = Alignment(horizontal="left")
    _largeurs(ws, [16, 12, 52])

    # ══ GAMMES ══ en-tête ligne 4, données à partir de la ligne 5 ═════════════
    ws = wb.create_sheet(FEUILLE_GAMMES)
    ws["A1"] = "Gammes — la suite ordonnée des opérations de chaque pièce"
    ws["A1"].font = _TITRE
    ws["A2"] = ("Une ligne par opération. 'pos' donne l'ordre de passage au sein "
                "de la pièce : l'opération pos=2 ne peut commencer qu'après la fin "
                "de l'opération pos=1.")
    ws["A2"].font = _AIDE
    ws["A3"] = ("Les colonnes OF et Désignation sont facultatives : elles servent "
                "uniquement à l'affichage du Gantt et des KPI.")
    ws["A3"].font = _AIDE
    _ecrire_entete(ws, 4, ["op_id", "job_id", "pos", "OF", "Désignation pièce"])
    _ecrire_lignes(ws, 5, [
        [1, 1, 1, "OF-1001", "Carter aluminium"],
        [2, 1, 2, "OF-1001", "Carter aluminium"],
        [3, 2, 1, "OF-1002", "Arbre de transmission"],
        [4, 2, 2, "OF-1002", "Arbre de transmission"],
        [5, 3, 1, "OF-1003", "Bride de fixation"],
        [6, 3, 2, "OF-1003", "Bride de fixation"],
    ])
    _largeurs(ws, [10, 10, 8, 14, 26])

    # ══ MODES & PT ══ en-tête ligne 4 ════════════════════════════════════════
    ws = wb.create_sheet(FEUILLE_MODES)
    ws["A1"] = "Modes & temps opératoires — quelle opération sur quelle machine"
    ws["A1"].font = _TITRE
    ws["A2"] = ("Une ligne par couple (opération, machine possible). Une opération "
                "qui peut tourner sur trois machines occupe trois lignes : "
                "le solveur choisira la meilleure.")
    ws["A2"].font = _AIDE
    ws["A3"] = "Chaque op_id de la feuille GAMMES doit apparaître au moins une fois ici."
    ws["A3"].font = _AIDE
    _ecrire_entete(ws, 4, ["op_id", "machine_id", "duree (min)"])
    _ecrire_lignes(ws, 5, [
        [1, 1, 45], [1, 2, 50],
        [2, 3, 30],
        [3, 1, 60], [3, 2, 55],
        [4, 4, 40],
        [5, 2, 35], [5, 3, 38],
        [6, 4, 25],
    ])
    _largeurs(ws, [10, 14, 14])

    # ══ TECHNICIENS ══ en-tête ligne 4 ═══════════════════════════════════════
    ws = wb.create_sheet(FEUILLE_TECHNICIENS)
    ws["A1"] = "Techniciens — qui fait le changement de série sur quelle machine"
    ws["A1"].font = _TITRE
    ws["A2"] = ("Un technicien peut gérer PLUSIEURS machines. En revanche chaque "
                "machine ne doit avoir QU'UN SEUL technicien : les parcs ne se "
                "recoupent pas.")
    ws["A2"].font = _AIDE
    ws["A3"] = ("Toute machine citée dans MODES & PT doit apparaître ici, sans quoi "
                "le fichier sera refusé.")
    ws["A3"].font = _AIDE
    _ecrire_entete(ws, 4, ["tech_id", "machine_id"])
    _ecrire_lignes(ws, 5, [
        [1, 1],
        [1, 2],
        [2, 3],
        [2, 4],
    ])
    _largeurs(ws, [12, 14])

    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


if __name__ == "__main__":
    from pathlib import Path
    chemin = Path(__file__).resolve().parent.parent / "template_gantt.xlsx"
    chemin.write_bytes(construire_template())
    print(f"Modèle écrit dans {chemin}")
